#!/usr/bin/env python3
"""
build_briefs.py -- generates assets/briefs.json

Sources:
  - briefs-src/**/*.docx                         seat brief documents (ID from filename only)
  - briefs-src/**/Seats_updated_Oct2026.xlsx      seat metadata (sheets Gp1-Seats (updated),
                                                   Gp2-Seats (updated); header row is row 2)

Output: assets/briefs.json, keyed by canonical seat ID, e.g.:
  {
    "AZ-01": {
      "type": "House", "title": "...", "status": "...", "dem": "...", "rep": "...",
      "others": "...", "ratings": "...", "notes": "...", "india_angle": "...",
      "india": "<html of the Significance-for-India bullet>",
      "sections": [{"heading": "Seat Overview", "html": "<ul>...</ul>"}, ...],
      "built": "YYYY-MM-DD"
    }, ...
  }

Never hand-edit assets/briefs.json -- this script generates it.

ID rules mirror assets/js/seat-ids.js exactly (toSeatId). If you change the
rules in one place, change them in the other.
"""

import glob
import os
import re
import sys
from datetime import date

try:
    import docx
except ImportError:
    sys.exit("Missing dependency: run `pip3 install python-docx` and retry.")

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency: run `pip3 install openpyxl` and retry.")


REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BRIEFS_SRC = os.path.join(REPO_ROOT, "briefs-src")
DATA_JSON_PATH = os.path.join(REPO_ROOT, "data.json")
OUTPUT_PATH = os.path.join(REPO_ROOT, "assets", "briefs.json")
SPREADSHEET_NAME = "Seats_updated_Oct2026.xlsx"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


# ─────────────────────────────────────────────────────────────────────────────
# Canonical seat ID rules -- mirror of assets/js/seat-ids.js. Keep both in sync.
# ─────────────────────────────────────────────────────────────────────────────

STATE_TO_ABBR = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR", "California": "CA",
    "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE", "Florida": "FL", "Georgia": "GA",
    "Hawaii": "HI", "Idaho": "ID", "Illinois": "IL", "Indiana": "IN", "Iowa": "IA",
    "Kansas": "KS", "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS", "Missouri": "MO",
    "Montana": "MT", "Nebraska": "NE", "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ",
    "New Mexico": "NM", "New York": "NY", "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH",
    "Oklahoma": "OK", "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT", "Vermont": "VT",
    "Virginia": "VA", "Washington": "WA", "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
    "District of Columbia": "DC",
}
ABBR_SET = set(STATE_TO_ABBR.values())
ABBR_TO_STATE = {v: k for k, v in STATE_TO_ABBR.items()}
NAME_LOWER_TO_ABBR = {k.lower(): v for k, v in STATE_TO_ABBR.items()}


def resolve_state_abbr(token):
    if not token:
        return None
    t = token.strip()
    up = t.upper()
    if up in ABBR_SET:
        return up
    return NAME_LOWER_TO_ABBR.get(t.lower())


def to_seat_id(kind, raw):
    """Python port of assets/js/seat-ids.js toSeatId(type, raw)."""
    if not kind or raw is None:
        return None
    kind = str(kind).strip().lower()
    s = str(raw).strip()
    if not s:
        return None

    if kind in ("senate", "governor"):
        abbr = resolve_state_abbr(s)
        if not abbr:
            return None
        return f"Senate-{abbr}" if kind == "senate" else f"Gov-{abbr}"

    if kind != "house":
        return None

    s = re.sub(r"\bdistrict\b|\bdist\.?\b", " ", s, flags=re.IGNORECASE)
    s = re.sub(r"[_|]", "-", s)
    s = re.sub(r"\s+", " ", s).strip()

    m = re.match(r"^(.+?)[\s-]+(AL|AT-?LARGE)$", s, flags=re.IGNORECASE)
    if m:
        abbr = resolve_state_abbr(m.group(1).strip())
        return f"{abbr}-AL" if abbr else None

    m = re.match(r"^(.+?)[\s-]+(\d{1,2})$", s)
    if m:
        abbr = resolve_state_abbr(m.group(1).strip())
        if not abbr:
            return None
        n = int(m.group(2))
        if n == 0:
            return f"{abbr}-AL"
        return f"{abbr}-{n:02d}"

    return None


def seat_title(kind, canonical_id, raw_seat):
    """Human-readable title for the brief, independent of inconsistent docx title lines."""
    kind = kind.lower()
    abbr = canonical_id.split("-")[-1] if kind != "house" else canonical_id.split("-")[0]
    state_name = ABBR_TO_STATE.get(abbr, abbr)
    if kind == "senate":
        return f"{state_name} (Senate)"
    if kind == "governor":
        return f"{state_name} (Governor)"
    if canonical_id.endswith("-AL"):
        return f"{state_name} At-Large (House)"
    district_num = int(canonical_id.split("-")[1])
    return f"{state_name} District {district_num} (House)"


# ─────────────────────────────────────────────────────────────────────────────
# .docx parsing
# ─────────────────────────────────────────────────────────────────────────────

FILENAME_RE = re.compile(r"^SB - (.+?) (House|Senate|Governor)\.docx$", re.IGNORECASE)
SECTION_RE = re.compile(r"^(\d+)\.\s*(.*)$")
SECTION_HEADINGS = {
    1: "Seat Overview",
    2: "Electoral Context",
    3: "Key Seat Issues",
    4: "Strategic Outlook",
    5: "Key Local Media",
}
INDIA_LABEL_RE = re.compile(r"significance for india", re.IGNORECASE)


def html_escape(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def run_is_bold(r_el):
    rpr = r_el.find(W_NS + "rPr")
    if rpr is None:
        return False
    b = rpr.find(W_NS + "b")
    if b is None:
        return False
    val = b.get(W_NS + "val")
    return val not in ("0", "false", "off")


def run_element_to_html(r_el):
    text = "".join(t.text or "" for t in r_el.findall(W_NS + "t"))
    if not text:
        return ""
    text = html_escape(text)
    return f"<strong>{text}</strong>" if run_is_bold(r_el) else text


def paragraph_to_inline_html(p_element, part):
    """Walk a <w:p>'s direct children in document order, turning runs into
    (optionally bold) text and <w:hyperlink> elements into real <a> tags."""
    pieces = []
    for child in p_element:
        tag = child.tag
        if tag == W_NS + "r":
            pieces.append(run_element_to_html(child))
        elif tag == W_NS + "hyperlink":
            rid = child.get(R_NS + "id")
            url = None
            if rid:
                rel = part.rels.get(rid)
                if rel is not None and rel.is_external:
                    url = rel.target_ref
            inner = "".join(run_element_to_html(r) for r in child.findall(W_NS + "r"))
            if url:
                pieces.append(f'<a href="{html_escape(url)}" target="_blank" rel="noopener">{inner}</a>')
            else:
                pieces.append(inner)
    return "".join(pieces)


def paragraph_text(p_element):
    return "".join(t.text or "" for t in p_element.iter(W_NS + "t")).strip()


class BriefParseError(Exception):
    pass


def parse_brief_docx(path):
    """Returns (sections, india_html). Raises BriefParseError on malformed input."""
    document = docx.Document(path)
    part = document.part
    paragraphs = document.paragraphs

    if len(paragraphs) < 3:
        raise BriefParseError("fewer than 3 paragraphs (expected title, 'Seat Brief', sections)")

    body = paragraphs[2:]  # drop title line + "Seat Brief" line

    # Locate section header paragraphs (no bullet numbering, matches "N. Heading").
    header_positions = []  # list of (index_in_body, section_number)
    for i, p in enumerate(body):
        txt = paragraph_text(p._p)
        m = SECTION_RE.match(txt)
        if not m:
            continue
        has_numpr = p._p.pPr is not None and p._p.pPr.numPr is not None
        if not has_numpr:
            header_positions.append((i, int(m.group(1))))

    expected_numbers = [1, 2, 3, 4, 5]
    found_numbers = [n for _, n in header_positions]
    if found_numbers != expected_numbers:
        raise BriefParseError(f"expected section headers 1..5 in order, found {found_numbers}")

    sections = []
    india_html = None
    for idx, (pos, num) in enumerate(header_positions):
        start = pos + 1
        end = header_positions[idx + 1][0] if idx + 1 < len(header_positions) else len(body)
        bullet_items = []
        for p in body[start:end]:
            txt = paragraph_text(p._p)
            if not txt:
                continue
            item_html = paragraph_to_inline_html(p._p, part)
            bullet_items.append(item_html)
            if INDIA_LABEL_RE.search(txt):
                india_html = item_html
        html = "<ul>" + "".join(f"<li>{item}</li>" for item in bullet_items) + "</ul>"
        sections.append({"heading": SECTION_HEADINGS[num], "html": html})

    return sections, india_html


def find_docx_files(root):
    return sorted(glob.glob(os.path.join(root, "**", "*.docx"), recursive=True))


FOLDER_CHAMBER_KEYWORDS = (
    ("Governor", "governor"),  # check before "senate"/"house" -- "Governors" only matches this one
    ("Senate", "senate"),
    ("House", "house"),
)


def folder_chamber(path):
    """Best-guess chamber implied by the immediate parent folder name, or None
    if the folder name doesn't clearly say. Never used for ID purposes --
    only to flag mismatches against the filename, which remains the source
    of truth for the seat ID."""
    parent = os.path.basename(os.path.dirname(path)).lower()
    for label, keyword in FOLDER_CHAMBER_KEYWORDS:
        if keyword in parent:
            return label
    return None


def build_briefs_from_docx(root, report):
    """Returns dict canonical_id -> {type, title, sections, india, built}."""
    briefs = {}
    id_sources = {}  # canonical_id -> source file path (to detect duplicates)

    for path in find_docx_files(root):
        rel_path = os.path.relpath(path, REPO_ROOT)
        filename = os.path.basename(path)

        if filename.startswith("~$"):
            continue  # Word lock file

        m = FILENAME_RE.match(filename)
        if not m:
            report["failed"].append((rel_path, "filename doesn't match 'SB - <seat> <House|Senate|Governor>.docx'"))
            continue

        raw_seat, chamber = m.group(1).strip(), m.group(2)

        f_chamber = folder_chamber(path)
        if f_chamber and f_chamber.lower() != chamber.lower():
            report["folder_mismatches"].append((rel_path, chamber.capitalize(), f_chamber))

        canonical_id = to_seat_id(chamber, raw_seat)
        if not canonical_id:
            report["failed"].append((rel_path, f"could not resolve seat ID from {chamber!r} + {raw_seat!r}"))
            continue

        if canonical_id in id_sources:
            report["duplicates"].setdefault(canonical_id, [id_sources[canonical_id]]).append(rel_path)
            continue
        id_sources[canonical_id] = rel_path

        try:
            sections, india_html = parse_brief_docx(path)
        except BriefParseError as e:
            report["failed"].append((rel_path, str(e)))
            del id_sources[canonical_id]
            continue
        except Exception as e:  # defensive: don't let one bad file kill the run
            report["failed"].append((rel_path, f"unexpected error: {e}"))
            del id_sources[canonical_id]
            continue

        briefs[canonical_id] = {
            "type": chamber.capitalize(),
            "title": seat_title(chamber, canonical_id, raw_seat),
            "sections": sections,
            "india": india_html,
            "built": date.today().isoformat(),
        }

    return briefs


# ─────────────────────────────────────────────────────────────────────────────
# Spreadsheet parsing
# ─────────────────────────────────────────────────────────────────────────────

FIELD_HEADER_CANDIDATES = {
    "type": "type",
    "seat_raw": "seat (as in original sheet)",
    "current_district": "current district",
    "status": "current status",
    "dem": "democratic nominee",
    "rep": "republican nominee",
    "others": "others on ballot",
    "ratings": "current forecaster ratings",
    "notes": "key notes",
    "india_angle": "india angle",
}

HOUSE_DISTRICT_PREFIX_RE = re.compile(r"^\s*([A-Za-z]{2}-\d{1,2})\b")


def find_spreadsheet(root):
    matches = glob.glob(os.path.join(root, "**", SPREADSHEET_NAME), recursive=True)
    return matches[0] if matches else None


def build_header_index(header_row):
    """Maps our field names -> 0-based column index, for one sheet's header row."""
    lower_headers = [(str(h).strip().lower() if h else "") for h in header_row]
    index = {}
    for field, candidate in FIELD_HEADER_CANDIDATES.items():
        found = None
        for i, h in enumerate(lower_headers):
            if candidate in h:
                found = i
                break
        index[field] = found
    return index


def cell(row, index, field):
    i = index.get(field)
    if i is None or i >= len(row):
        return None
    v = row[i]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def clean_house_cell(raw):
    """Strips trailing parenthetical notes, e.g. 'FL-14 (redrawn, ~Trump+10)' -> 'FL-14'."""
    if not raw:
        return raw
    m = HOUSE_DISTRICT_PREFIX_RE.match(str(raw))
    return m.group(1) if m else raw


def build_rows_from_sheet(xlsx_path, report):
    """Returns dict canonical_id -> row data dict."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    rows_by_id = {}

    for sheet_name in wb.sheetnames:
        if "(updated)" not in sheet_name:
            continue  # skip "Gp1 - Original sheet" / "Gp2-Original" raw tabs
        ws = wb[sheet_name]
        header_row = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
        index = build_header_index(header_row)

        for r in range(3, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            chamber = cell(row, index, "type")
            if chamber not in ("Senate", "House", "Governor"):
                continue  # footer/legend/junk row

            if chamber == "House":
                current_district = cell(row, index, "current_district")
                seat_raw = clean_house_cell(current_district) if current_district else cell(row, index, "seat_raw")
            else:
                seat_raw = cell(row, index, "seat_raw")

            canonical_id = to_seat_id(chamber, seat_raw)
            if not canonical_id:
                report["sheet_unparsed"].append((sheet_name, r, chamber, seat_raw))
                continue

            data = {
                "type": chamber,
                "status": cell(row, index, "status"),
                "dem": cell(row, index, "dem"),
                "rep": cell(row, index, "rep"),
                "others": cell(row, index, "others"),
                "ratings": cell(row, index, "ratings"),
                "notes": cell(row, index, "notes"),
                "india_angle": cell(row, index, "india_angle"),
            }

            if canonical_id in rows_by_id:
                report["sheet_duplicates"].setdefault(canonical_id, [rows_by_id[canonical_id][1]]).append(
                    f"{sheet_name}!row{r}"
                )
                continue
            rows_by_id[canonical_id] = (data, f"{sheet_name}!row{r}")

    return {k: v[0] for k, v in rows_by_id.items()}


# ─────────────────────────────────────────────────────────────────────────────
# Site-seat comparison (data.json)
# ─────────────────────────────────────────────────────────────────────────────

def site_seat_ids():
    import json

    if not os.path.exists(DATA_JSON_PATH):
        return set()
    with open(DATA_JSON_PATH, encoding="utf-8") as f:
        data = json.load(f)

    ids = set()
    for race in data.get("races", {}).get("senate", {}).get("races", []):
        sid = to_seat_id("senate", race.get("state"))
        if sid:
            ids.add(sid)
    for race in data.get("races", {}).get("house", {}).get("competitive", []):
        sid = to_seat_id("house", race.get("district"))
        if sid:
            ids.add(sid)
    for race in data.get("races", {}).get("governor", {}).get("races", []):
        sid = to_seat_id("governor", race.get("state"))
        if sid:
            ids.add(sid)
    return ids


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    import json

    if not os.path.isdir(BRIEFS_SRC):
        sys.exit(f"briefs-src/ not found at {BRIEFS_SRC}")

    report = {
        "failed": [],              # [(rel_path, reason)]
        "duplicates": {},          # id -> [rel_paths]
        "folder_mismatches": [],   # [(rel_path, filename_chamber, folder_chamber)]
        "sheet_unparsed": [],      # [(sheet, row, type, raw)]
        "sheet_duplicates": {},    # id -> [sheet!row]
    }

    docx_briefs = build_briefs_from_docx(BRIEFS_SRC, report)

    xlsx_path = find_spreadsheet(BRIEFS_SRC)
    sheet_rows = {}
    if not xlsx_path:
        print(f"WARNING: {SPREADSHEET_NAME} not found under {BRIEFS_SRC} -- output will have no metadata fields.")
    else:
        sheet_rows = build_rows_from_sheet(xlsx_path, report)

    all_ids = set(docx_briefs) | set(sheet_rows)
    output = {}
    for seat_id in sorted(all_ids):
        entry = dict(docx_briefs.get(seat_id, {}))
        meta = sheet_rows.get(seat_id, {})
        for field in ("type", "status", "dem", "rep", "others", "ratings", "notes", "india_angle"):
            entry.setdefault(field, None)
            if meta.get(field) is not None:
                entry[field] = meta[field]
        if "type" not in entry or entry.get("type") is None:
            entry["type"] = meta.get("type")
        output[seat_id] = entry

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
        f.write("\n")

    briefs_with_no_sheet_row = sorted(set(docx_briefs) - set(sheet_rows))
    sheet_rows_with_no_brief = sorted(set(sheet_rows) - set(docx_briefs))
    site_ids = site_seat_ids()
    site_seats_with_no_brief = sorted(site_ids - set(docx_briefs))

    # ── Report ──────────────────────────────────────────────────────────────
    print("=" * 70)
    print("BUILD REPORT")
    print("=" * 70)
    print(f"\nBriefs built: {len(output)}  (written to {os.path.relpath(OUTPUT_PATH, REPO_ROOT)})")

    print(f"\nFile names that failed to parse: {len(report['failed'])}")
    for rel_path, reason in report["failed"]:
        print(f"  - {rel_path}: {reason}")

    print(f"\nDuplicate IDs (docx): {len(report['duplicates'])}")
    for seat_id, paths in report["duplicates"].items():
        print(f"  - {seat_id}: {paths}")

    print(f"\nFolder mismatches (filename says one chamber, parent folder says another): {len(report['folder_mismatches'])}")
    for rel_path, name_chamber, dir_chamber in report["folder_mismatches"]:
        print(f"  - {rel_path}: filename says {name_chamber}, folder says {dir_chamber} (seat ID used the filename)")

    print(f"\nDuplicate IDs (spreadsheet rows): {len(report['sheet_duplicates'])}")
    for seat_id, locs in report["sheet_duplicates"].items():
        print(f"  - {seat_id}: {locs}")

    print(f"\nSpreadsheet rows that could not be resolved to an ID: {len(report['sheet_unparsed'])}")
    for sheet_name, r, chamber, raw in report["sheet_unparsed"]:
        print(f"  - {sheet_name}!row{r} ({chamber}, seat={raw!r})")

    print(f"\nSheet rows with no matching brief: {len(sheet_rows_with_no_brief)}")
    for seat_id in sheet_rows_with_no_brief:
        print(f"  - {seat_id}")

    print(f"\nBriefs with no matching sheet row: {len(briefs_with_no_sheet_row)}")
    for seat_id in briefs_with_no_sheet_row:
        print(f"  - {seat_id}")

    print(f"\nSite seats (data.json) with no brief: {len(site_seats_with_no_brief)}")
    for seat_id in site_seats_with_no_brief:
        print(f"  - {seat_id}")

    print("\n" + "=" * 70)


if __name__ == "__main__":
    main()
